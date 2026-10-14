from __future__ import annotations

import re
from html.parser import HTMLParser
from pathlib import Path


class _ReadableHtml(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.ignored = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "noscript"}:
            self.ignored += 1
        if not self.ignored and tag in {"p", "div", "br", "tr", "li", "h1", "h2", "h3"}:
            self.parts.append("\n\n")

    def handle_endtag(self, tag):
        if tag in {"script", "style", "noscript"} and self.ignored:
            self.ignored -= 1
        if not self.ignored and tag in {"td", "th"}:
            self.parts.append(" | ")

    def handle_data(self, data):
        if not self.ignored:
            self.parts.append(data)


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".csv"}:
        return path.read_text(encoding="utf-8-sig")
    if suffix in {".html", ".htm"}:
        try:
            html = path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError:
            html = path.read_text(encoding="gb18030")
        parser = _ReadableHtml()
        parser.feed(html)
        text = normalize_text("".join(parser.parts))
        if len(text) < 80:
            raise ValueError("HTML正文过短，可能是错误页或需登录页面，请核验原件")
        return text
    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("读取DOCX需要安装python-docx") from exc
        doc = Document(path)
        paragraphs = [paragraph.text for paragraph in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(paragraphs)
    if suffix == ".pdf":
        try:
            import pymupdf as fitz
        except ImportError as exc:
            raise RuntimeError("读取PDF需要安装pymupdf") from exc
        with fitz.open(path) as pdf:
            text = "\n".join(page.get_text("text") for page in pdf)
        substantive = re.sub(r"住房城乡建设部信息公开|浏览专用|\s+", "", text)
        if len(substantive) < 80:
            raise ValueError(f"{path.name}可能是扫描件，请先OCR并人工校对后保存为UTF-8 Markdown")
        return text
    raise ValueError(f"暂不支持的知识库文件格式：{suffix}")


def normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def chunk_text(text: str, *, chunk_size: int = 1200, overlap: int = 180) -> list[str]:
    if type(chunk_size) is not int or type(overlap) is not int or chunk_size <= 0 or not 0 <= overlap < chunk_size:
        raise ValueError("分块长度须为正整数，重叠须为小于分块长度的非负整数")
    text = normalize_text(text)
    if not text:
        return []
    paragraphs = [part.strip() for part in text.split("\n\n") if part.strip()]
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        if len(paragraph) > chunk_size:
            if current:
                chunks.append(current)
                current = ""
            start = 0
            while start < len(paragraph):
                chunks.append(paragraph[start : start + chunk_size])
                start += max(1, chunk_size - overlap)
            continue
        candidate = f"{current}\n\n{paragraph}".strip()
        if current and len(candidate) > chunk_size:
            chunks.append(current)
            tail_length = min(overlap, max(0, chunk_size - len(paragraph) - 2))
            tail = current[-tail_length:] if tail_length else ""
            current = f"{tail}\n\n{paragraph}".strip()
        else:
            current = candidate
    if current:
        chunks.append(current)
    return chunks
